from pathlib import Path

from docx import Document as DocxDocument

from backend.ingestion.structure.document_model import (
    Document,
    Page,
    RawBlock,
)


def create_document_id(file_path: str) -> str:
    filename = Path(file_path).stem.lower()

    return "".join(
        char if char.isalnum() else "_"
        for char in filename
    ).strip("_")


def extract_docx(file_path: str) -> Document:

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"File not found: {file_path}"
        )

    if path.suffix.lower() != ".docx":
        raise ValueError(
            "Expected a DOCX file."
        )

    source_document = DocxDocument(file_path)

    blocks = []
    tables = []

    block_counter = 0

    # Paragraphs
    for paragraph in source_document.paragraphs:

        text = paragraph.text.strip()

        if not text:
            continue

        block_counter += 1

        blocks.append(
            RawBlock(
                block_id=f"b{block_counter}",
                page_number=1,
                text=text,
                block_type="paragraph",
                metadata={
                    "style": paragraph.style.name
                    if paragraph.style
                    else None
                },
            )
        )

    # Tables
    for table_index, table in enumerate(
        source_document.tables,
        start=1
    ):
        table_rows = []
        table_headers = []

        for row_index, row in enumerate(
            table.rows,
            start=1
        ):

            values = [
                cell.text.strip()
                for cell in row.cells
            ]

            text = " | ".join(values)
            if row_index == 1:
                table_headers = values
            else:
                table_rows.append(dict(zip(table_headers or [f"column_{i+1}" for i in range(len(values))], values)))

            if not text:
                continue

            block_counter += 1

            blocks.append(
                RawBlock(
                    block_id=f"table{table_index}_row{row_index}",
                    page_number=1,
                    text=text,
                    block_type="table_row",
                    metadata={
                        "table_index": table_index,
                        "row_index": row_index,
                        "columns": values,
                    },
                )
            )

        tables.append({
            "table_id": f"docx_table_{table_index}",
            "columns": table_headers,
            "rows": table_rows,
        })

    page = Page(
        page_number=1,
        blocks=blocks,
    )

    return Document(
        document_id=create_document_id(file_path),
        file_name=path.name,
        file_type="docx",
        title=path.stem,
        tables=tables,
        pages=[page],
        metadata={
            "parser": "python-docx"
        },
    )